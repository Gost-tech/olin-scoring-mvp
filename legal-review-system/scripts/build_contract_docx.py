#!/usr/bin/env python3
"""Build a review-friendly DOCX from the authoritative Olin contract Markdown."""

from __future__ import annotations

from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "244C66"
MUTED = "65727C"
RED = "9B1C1C"
LIGHT_YELLOW = "FFF2CC"
BODY = "222222"
PLACEHOLDER_RE = re.compile(r"(\[[^\]]+\]|\[●\]|\[•\])")


def set_font(run, *, name="Calibri", size=11, color=BODY, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_run(run, fill=LIGHT_YELLOW):
    rpr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    rpr.append(shd)


def add_rich_text(paragraph, text, *, size=11, color=BODY, italic=False):
    cursor = 0
    bold = False
    token_re = re.compile(r"(\*\*|`[^`]+`|\[[^\]]+\])")
    for match in token_re.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_font(run, size=size, color=color, bold=bold, italic=italic)
        token = match.group(0)
        if token == "**":
            bold = not bold
        else:
            literal = token[1:-1] if token.startswith("`") else token
            run = paragraph.add_run(literal)
            set_font(run, size=size, color=color, bold=bold, italic=italic)
            if token.startswith("["):
                shade_run(run)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_font(run, size=size, color=color, bold=bold, italic=italic)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    set_font(run, size=9, color=MUTED)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (16, 14, 8, BLUE),
        "Heading 2": (13, 11, 6, BLUE),
        "Heading 3": (12, 8, 4, BLUE),
    }
    for name, (size, before, after, color) in specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(11)
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.188)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.25


def configure_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("OLIN  |  ACUERDO DE PILOTO SOMBRA  |  BORRADOR PARA REVISIÓN LEGAL")
    set_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Confidencial · No firmar · Página ")
    set_font(run, size=9, color=MUTED)
    add_page_field(footer)


def build(source: Path, output: Path):
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    doc.core_properties.title = "Acuerdo de piloto sombra Olin - Borrador para revisión legal"
    doc.core_properties.subject = "Borrador no firmable para revisión por abogado mexicano"
    doc.core_properties.author = "Olin"

    lines = source.read_text(encoding="utf-8").splitlines()
    first_title = True
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line == "---":
            continue
        if line.startswith("# "):
            if first_title:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(20)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run(line[2:])
                set_font(run, size=20, color=RED, bold=True)
                first_title = False
            else:
                p = doc.add_paragraph(line[2:], style="Heading 1")
            continue
        if line.startswith("## "):
            text = line[3:]
            if text.lower().startswith("acuerdo de piloto"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(18)
                run = p.add_run(text)
                set_font(run, size=15, color=BLUE, bold=True)
            else:
                doc.add_paragraph(text, style="Heading 1")
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, line[2:])
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Inches(0.375)
            p.paragraph_format.first_line_indent = Inches(-0.188)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            add_rich_text(p, re.sub(r"^\d+\.\s+", "", line))
            continue
        p = doc.add_paragraph()
        if re.match(r"^\d+\.\d+\s", line):
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.35)
            p.paragraph_format.keep_together = True
        add_rich_text(p, line)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print("usage: build_contract_docx.py SOURCE.md OUTPUT.docx", file=sys.stderr)
        return 2
    build(Path(argv[1]).resolve(), Path(argv[2]).resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
