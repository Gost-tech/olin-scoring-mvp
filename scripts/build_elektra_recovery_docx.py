from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


INK = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(90, 99, 110)
WHITE = "FFFFFF"
LIGHT = "F2F4F7"
CALLOUT = "E8EEF5"
RISK = RGBColor(155, 28, 28)


def font(run, size=None, bold=None, color=None, name="Calibri", italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, INK, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("OLIN × ELEKTRA  |  TECHNICAL EVALUATION")
    font(run, size=8.5, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def clean_inline(text):
    text = re.sub(r"\[([^]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("`", "")
    return text


def add_rich_para(doc, text, style=None, bold_prefix=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.widow_control = True
    text = clean_inline(text)
    if bold_prefix and ":" in text:
        head, tail = text.split(":", 1)
        r = p.add_run(head + ":")
        font(r, bold=True)
        r = p.add_run(tail)
        font(r)
    else:
        r = p.add_run(text)
        font(r)
    return p


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(clean_inline(text))
    font(r, size=11, bold=True, color=INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_markdown_table(doc, rows):
    if len(rows) < 2:
        return
    data = rows[1:] if all(re.fullmatch(r":?-+:?", c.strip()) for c in rows[1]) else rows
    cols = max(len(r) for r in data)
    table = doc.add_table(rows=len(data), cols=cols)
    table.style = "Table Grid"
    widths = [9360 // cols] * cols
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    for i, row in enumerate(data):
        for j in range(cols):
            cell = table.cell(i, j)
            if i == 0:
                shade(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            value = clean_inline(row[j].strip()) if j < len(row) else ""
            r = p.add_run(value)
            font(r, size=8.5, bold=(i == 0), color=INK if i == 0 else None)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build(source: Path, output: Path):
    doc = Document()
    configure(doc)
    lines = source.read_text(encoding="utf-8").splitlines()

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("OLIN × ELEKTRA")
    font(r, size=24, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    r = subtitle.add_run("Seven-Day Technical Recovery & Evaluation Plan")
    font(r, size=15, bold=True, color=BLUE)
    for label, value in (
        ("Decision date", "1 September 2026"),
        ("Target", "Sandbox technical-evaluation build in seven calendar days"),
        ("Audience", "Olin engineering and Elektra / Banco Azteca reviewers"),
        ("Classification", "Evaluation planning — no production approval implied"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label + ": ")
        font(r, size=10, bold=True, color=INK)
        r = p.add_run(value)
        font(r, size=10, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    in_code = False
    code_lines = []
    table_rows = []
    skipped_title = False
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.space_after = Pt(8)
                r = p.add_run("\n".join(code_lines))
                font(r, size=8.3, name="Menlo", color=INK)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("|") and line.endswith("|"):
            table_rows.append([c for c in line.strip("|").split("|")])
            continue
        if table_rows:
            add_markdown_table(doc, table_rows)
            table_rows = []
        if not line:
            continue
        if line.startswith("# "):
            if not skipped_title:
                skipped_title = True
                continue
            add_rich_para(doc, line[2:], "Heading 1")
        elif line.startswith("## "):
            add_rich_para(doc, line[3:], "Heading 1")
        elif line.startswith("### "):
            add_rich_para(doc, line[4:], "Heading 2")
        elif line.startswith("#### "):
            add_rich_para(doc, line[5:], "Heading 3")
        elif line.startswith("> "):
            add_callout(doc, line[2:])
        elif re.match(r"^- \[[ xX]\] ", line):
            add_rich_para(doc, "☐ " + line[6:], "List Bullet")
        elif line.startswith("- "):
            add_rich_para(doc, line[2:], "List Bullet", bold_prefix=True)
        elif re.match(r"^\d+\. ", line):
            add_rich_para(doc, re.sub(r"^\d+\. ", "", line), "List Number")
        elif line.startswith("**") and ":**" in line:
            add_rich_para(doc, line, bold_prefix=True)
        else:
            p = add_rich_para(doc, line)
            if line.startswith("Any failed item") or line.startswith("The release is ready"):
                for run in p.runs:
                    run.font.color.rgb = RISK
                    run.bold = True
    if table_rows:
        add_markdown_table(doc, table_rows)

    props = doc.core_properties
    props.title = "Olin × Elektra — Seven-Day Technical Recovery & Evaluation Plan"
    props.subject = "Architecture, execution, model governance and acceptance gates"
    props.author = "Olin"
    props.keywords = "Olin, Elektra, Banco Azteca, architecture, credit, model risk"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_elektra_recovery_docx.py SOURCE.md OUTPUT.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
